from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Trovun Marketplace Policies and Community Standards.docx"
LOGO = ROOT / "Trovun Logo No Background.png"

INK = "111720"
NAVY = "101923"
GOLD = "E7BC35"
PALE_GOLD = "FBF4D9"
PALE_BLUE = "EEF3F7"
PALE_GRAY = "F5F5F4"
MID_GRAY = "666666"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), LINE)


def set_keep(paragraph, *, next_=False, together=False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if next_:
        p_pr.append(OxmlElement("w:keepNext"))
    if together:
        p_pr.append(OxmlElement("w:keepLines"))
    p_pr.append(OxmlElement("w:widowControl"))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, end])


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "385D7A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_para(doc, text: str = "", *, bold_lead: str | None = None, style=None, keep=False):
    paragraph = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    if keep:
        set_keep(paragraph, together=True)
    return paragraph


def add_bullets(doc, items, *, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        if isinstance(item, tuple):
            lead, rest = item
            p.add_run(lead).bold = True
            p.add_run(rest)
        else:
            p.add_run(item)
        set_keep(p, together=True)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        set_keep(p, together=True)


def add_policy_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, (cell, label, width) in enumerate(zip(header.cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx else WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        for idx, (cell, value, width) in enumerate(zip(row.cells, row_data, widths)):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else PALE_BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            r.font.size = Pt(9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text: str, level: int, *, page_break=False):
    p = doc.add_heading(text, level=level)
    if page_break:
        p.paragraph_format.page_break_before = True
    set_keep(p, next_=True, together=True)
    return p


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.title = "Trovun Marketplace Policies and Community Standards"
    doc.core_properties.subject = "Marketplace guidelines terms of use privacy and notifications"
    doc.core_properties.author = "Trovun"
    doc.core_properties.keywords = "Trovun marketplace policy privacy terms students Waterloo"

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(13.5)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(18)

    for style_name, size, before, after in (
        ("Heading 1", 20, 14, 8),
        ("Heading 2", 14.5, 12, 6),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for list_style in ("List Bullet", "List Bullet 2", "List Number"):
        styles[list_style].font.name = "Aptos"
        styles[list_style].font.size = Pt(10.5)
        styles[list_style].paragraph_format.space_after = Pt(4)

    # Cover
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.autofit = False
    cell = banner.cell(0, 0)
    set_cell_width(cell, 6.94)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=220, start=260, bottom=220, end=260)
    set_table_borders(banner)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(0.88))
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("T R O V U N")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(WHITE)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(34)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Trovun Marketplace Policies and Community Standards")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Terms of Use Privacy Statement and Marketplace Guidelines")

    add_para(doc, "Effective 21 September 2026", bold_lead="Effective ")
    add_para(doc, "Version 1.0 publication draft", bold_lead="Version ")
    add_para(
        doc,
        "Applies to the Trovun website, student marketplace, account features, listings, messages, ratings, notifications, and related services.",
        keep=True,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    p = add_para(
        doc,
        "Trovun is an independent student built marketplace for the University of Waterloo community. It is not affiliated with, endorsed by, or operated by the University of Waterloo.",
        keep=True,
    )
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(16)
    p = add_para(
        doc,
        "Publication note  This document is a comprehensive operational draft and is not legal advice. Before publication, Trovun should confirm the legal operator name, service address, monitored support and privacy contacts, minimum age rule, and final retention schedule with qualified Canadian counsel.",
        bold_lead="Publication note  ",
        keep=True,
    )
    p.runs[0].font.color.rgb = RGBColor.from_string(INK)
    doc.add_page_break()

    # Header and footer from page two onward.
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "TROVUN    MARKETPLACE POLICIES"
    hp.style = styles["Caption"]
    hp.runs[0].font.name = "Aptos"
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.bold = True
    hp.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "Effective 21 September 2026    Draft for legal review"
    fp.runs[0].font.name = "Aptos"
    fp.runs[0].font.size = Pt(8)
    fp.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_number(fp)

    add_heading(doc, "How to use this document", 1)
    add_para(
        doc,
        "This document combines the rules that students need to follow with the contractual and privacy terms that govern Trovun. The quick rules and marketplace guidelines are written for daily use. The Terms of Use, Privacy Statement, and Notification Policy provide the fuller legal and operational framework.",
    )
    add_para(
        doc,
        "If a short rule conflicts with a more detailed provision, the detailed provision controls. Applicable law always controls over this document.",
    )

    add_heading(doc, "Contents", 2)
    contents = [
        "1  Quick rules for every student",
        "2  Marketplace and community guidelines",
        "3  Items and services that cannot be listed",
        "4  Transaction and meetup safety",
        "5  Terms of Use",
        "6  Privacy Statement",
        "7  Notification Policy",
        "8  Reporting moderation and appeals",
        "9  Legal references and publication checklist",
    ]
    add_bullets(doc, contents)

    add_heading(doc, "Defined terms", 2)
    add_policy_table(
        doc,
        ["Term", "Meaning"],
        [
            ("Trovun", "The independent student marketplace service, including its website and related features."),
            ("User", "A person who accesses or uses Trovun, including a buyer, seller, or moderator."),
            ("Listing", "An offer posted by a user for an item, including its description, photos, price, condition, and meetup details."),
            ("Transaction", "An arrangement between users to inspect, exchange, buy, sell, or give away an item."),
            ("Content", "Listings, photos, profile information, messages, ratings, reports, and other material submitted through Trovun."),
            ("Service", "The Trovun website, marketplace, messaging, notification, moderation, and account features."),
        ],
        [1.4, 5.3],
    )

    # Part 1
    add_heading(doc, "1  Quick rules for every student", 1, page_break=True)
    add_para(
        doc,
        "Use Trovun for honest, lawful student to student exchanges. Keep all listing and account information accurate, protect personal information, meet in a safe public place, and stop the transaction if anything feels wrong.",
    )
    add_numbers(
        doc,
        [
            "Use your own verified Waterloo student account and keep it secure.",
            "List only items you own or are authorized to sell.",
            "Describe the item, price, condition, defects, included parts, and pickup terms accurately.",
            "Never list a prohibited, recalled, stolen, counterfeit, unsafe, regulated, or age restricted item.",
            "Keep conversations respectful and related to the transaction.",
            "Do not request passwords, one time codes, banking credentials, government identification numbers, or unnecessary personal information.",
            "Inspect the item before paying and use a payment method you understand.",
            "Meet in a familiar public place and tell someone where you are going.",
            "Mark completed listings as sold and rate only a seller with whom you completed a genuine transaction.",
            "Report suspected fraud, harassment, unsafe products, or other serious rule violations promptly.",
        ],
    )

    add_heading(doc, "What Trovun does", 2)
    add_para(
        doc,
        "Trovun provides account verification, listing tools, search, private messaging, seller ratings, notification controls, and moderation. Trovun may review reports, remove listings, restrict accounts, preserve records needed for safety or legal compliance, and cooperate with lawful authorities.",
    )
    add_heading(doc, "What Trovun does not do", 2)
    add_para(
        doc,
        "Trovun is not the buyer, seller, owner, manufacturer, appraiser, delivery provider, payment processor, insurer, or agent in a user transaction. Trovun does not take possession of items, set the final transaction price, guarantee payment, inspect every listing, or promise that another user will complete a transaction.",
    )

    # Part 2
    add_heading(doc, "2  Marketplace and community guidelines", 1, page_break=True)
    add_heading(doc, "Account eligibility and identity", 2)
    add_bullets(
        doc,
        [
            "Use Trovun only if you are eligible under the registration rules shown in the Service and can legally agree to these policies.",
            "Register with your own eligible University of Waterloo email address. Do not share, sell, transfer, or create an account for another person.",
            "Provide accurate profile information. A shortened display name and selected trust information may be shown to other verified users, while private account information remains restricted.",
            "Do not impersonate another person, misstate your student status, evade a suspension, or create multiple accounts to manipulate ratings or listings.",
            "Keep your password and verification codes confidential. Notify Trovun promptly if you believe your account has been compromised.",
        ],
    )

    add_heading(doc, "Listing standards", 2)
    add_para(doc, "Every listing must enable a reasonable student to understand what is being offered and make an informed decision.")
    add_bullets(
        doc,
        [
            ("Ownership and authority  ", "You must own the item or have clear authority to sell it."),
            ("Accurate title and category  ", "Use a specific title and the most appropriate marketplace category."),
            ("Complete description  ", "State the actual condition, material defects, missing parts, compatibility limits, modifications, recall history, and whether accessories are included."),
            ("Current photographs  ", "Use photographs of the actual item. Do not use misleading stock images, altered images that hide defects, or images copied without permission."),
            ("Clear price  ", "State the price in Canadian dollars and disclose whether offers are accepted. Do not use bait pricing, undisclosed mandatory charges, or artificial urgency."),
            ("Safe meetup details  ", "Choose a public meetup point. Do not publish private residence access codes, room numbers, schedules, or other sensitive location information."),
            ("One item one honest listing  ", "Do not flood the marketplace with duplicates, keyword spam, unrelated tags, or misleading brand names."),
            ("Prompt status updates  ", "Remove unavailable listings or mark them sold promptly."),
        ],
    )

    add_heading(doc, "Messaging standards", 2)
    add_bullets(
        doc,
        [
            "Use messaging to ask about the item, negotiate, arrange inspection, and agree on a safe public meetup.",
            "Do not send harassment, threats, discriminatory content, sexual content, repeated unwanted messages, spam, or mass solicitations.",
            "Do not send malicious links, malware, fraudulent payment requests, or messages designed to move a user into an unsafe or unverifiable transaction.",
            "Never ask for a password, authentication code, student login, Social Insurance Number, full banking credential, or copy of government identification.",
            "Do not publish another person’s message or personal information without a lawful basis or their permission.",
        ],
    )

    add_heading(doc, "Ratings and feedback", 2)
    add_para(
        doc,
        "A buyer may rate a seller through a conversation connected to a listing that the seller has marked sold. Ratings must reflect a genuine transaction and the user’s honest experience. Users must not trade ratings, threaten a negative rating to obtain an unrelated benefit, create accounts to manipulate scores, or submit discriminatory or retaliatory feedback. Trovun may remove or disregard ratings that appear fraudulent, abusive, or unrelated to a completed exchange.",
    )

    add_heading(doc, "Respectful conduct", 2)
    add_para(
        doc,
        "Treat other users with respect. Discrimination, harassment, stalking, intimidation, threats, hate speech, sexual exploitation, and conduct that creates a credible safety risk are prohibited. A disagreement over price does not justify abusive conduct. End the conversation and report the matter if another user crosses these boundaries.",
    )

    # Part 3
    add_heading(doc, "3  Items and services that cannot be listed", 1, page_break=True)
    add_para(
        doc,
        "The following categories are prohibited even if possession may be lawful in another context. This list is not exhaustive. Trovun may remove any listing that is unlawful, unsafe, unsuitable for a student marketplace, or likely to expose users or the Service to harm.",
    )
    prohibited_rows = [
        ("Illegal and stolen property", "Stolen goods, items obtained through fraud, burglary tools, counterfeit currency, government property offered without authority, or any item whose sale or possession is unlawful."),
        ("Weapons and explosives", "Firearms, ammunition, replicas presented as real, prohibited knives, explosives, fireworks, weapon components, or instructions and materials primarily intended to create a weapon."),
        ("Drugs and controlled substances", "Illegal drugs, controlled substances, drug paraphernalia, prescription medication, cannabis, cannabis products, or substances marketed for intoxication."),
        ("Alcohol tobacco and vaping", "Alcohol, tobacco, nicotine products, vaping devices, vape liquids, or related age restricted products."),
        ("Unsafe and recalled products", "Recalled, banned, non compliant, expired, contaminated, materially damaged, or hazardous products, including unsafe chargers, batteries, helmets, children’s products, and appliances."),
        ("Counterfeit and infringing goods", "Counterfeit branded goods, unauthorized replicas, pirated media, copied access codes, circumvention devices, or content sold without the necessary intellectual property rights."),
        ("Academic misconduct", "Completed assignments, exam questions obtained without authorization, answer keys, stolen course materials, essay writing, impersonation, credential sharing, or services intended to facilitate academic misconduct."),
        ("Accounts credentials and access", "Student accounts, software accounts, streaming accounts, passwords, authentication codes, identity documents, access cards, keys, transit credentials, or tickets that cannot lawfully be transferred."),
        ("Financial and regulated products", "Currency exchange, loans, credit, securities, cryptocurrency schemes, gift cards with unverifiable balances, lottery products, or other regulated financial instruments."),
        ("Food health and personal safety", "Homemade food offered commercially, opened consumables, unsafe supplements, medical devices requiring authorization, bodily fluids, human remains, or products making unapproved medical claims."),
        ("Animals and living materials", "Animals, pets, wildlife, invasive species, biological specimens, or other regulated living material."),
        ("Sexual exploitation and adult services", "Sexual services, exploitative content, intimate images shared without consent, or content involving minors."),
        ("Housing employment and unrelated services", "Rental deposits, sublets, jobs, tutoring services, commercial advertising, fundraising, or other offers outside the item marketplace unless Trovun expressly introduces a suitable category."),
        ("Personal information and surveillance", "Personal databases, class lists, doxxing material, hidden cameras, unlawful tracking devices, or information obtained without consent."),
    ]
    add_policy_table(doc, ["Category", "Examples and scope"], prohibited_rows, [1.9, 4.8])

    add_heading(doc, "Seller checks before posting", 2)
    add_bullets(
        doc,
        [
            "Check the Government of Canada recalls and safety alerts database when an item could have a safety recall.",
            "Confirm that serial numbers and identifying labels are intact and that the item is not activation locked or reported lost or stolen.",
            "Remove personal data from electronics and reset devices only after preserving anything you need.",
            "Disclose repairs, third party parts, damaged batteries, missing safety equipment, and known defects.",
            "Do not list an item if you are uncertain whether you may lawfully possess, transfer, or sell it.",
        ],
    )

    # Part 4
    add_heading(doc, "4  Transaction and meetup safety", 1, page_break=True)
    add_heading(doc, "Before agreeing to buy", 2)
    add_bullets(
        doc,
        [
            "Ask for the model number, condition, included accessories, proof of ownership when appropriate, and a clear explanation of defects.",
            "Compare the price with credible alternatives. An unusually low price, rushed deadline, or refusal to answer ordinary questions may signal fraud.",
            "Keep material terms in Trovun messages so both users have a record of what was agreed.",
            "Do not send deposits for an item you have not inspected unless you knowingly accept that risk and use a payment method with appropriate protection.",
        ],
    )

    add_heading(doc, "Meeting safely", 2)
    add_bullets(
        doc,
        [
            "Meet during reasonable hours in a familiar, well lit public location with other people nearby.",
            "Tell a friend where you are going, who you are meeting, and when you expect to return.",
            "Bring another person for high value or bulky items when practical.",
            "Do not enter a private residence or vehicle if you feel uncomfortable. Leave immediately if the item, location, or person differs materially from what was represented.",
            "Call emergency services when there is an immediate threat. Trovun reporting is not a substitute for emergency assistance.",
        ],
    )

    add_heading(doc, "Inspection and payment", 2)
    add_para(
        doc,
        "Buyers should inspect and, where safe, test the item before paying. Confirm that activation locks are removed, serial numbers are present, batteries are not swollen or damaged, and important functions work. Sellers should allow a reasonable inspection without surrendering control of the item before payment.",
    )
    add_para(
        doc,
        "Users choose and assume responsibility for their payment method. Verify that payment has actually settled; screenshots, pending transfers, and email notices may be forged. Trovun does not process payments, hold deposits, provide escrow, reverse transfers, issue refunds, or guarantee chargeback rights.",
    )

    add_heading(doc, "Returns disputes and records", 2)
    add_para(
        doc,
        "Users should agree on any return or refund terms before completing the exchange. Private sales may not include the same return rights as a retail purchase, while a user acting as a business seller may have additional legal duties. Trovun may preserve relevant records and assist with moderation, but it does not decide ownership disputes or compel payment, refunds, returns, or repairs.",
    )

    # Terms
    add_heading(doc, "5  Terms of Use", 1, page_break=True)
    add_heading(doc, "Acceptance and scope", 2)
    add_para(
        doc,
        "By creating an account, accessing the private marketplace, posting a listing, sending a message, submitting a rating, or otherwise using the Service, you agree to these Terms of Use and the other policies in this document. If you do not agree, do not use the Service.",
    )
    add_para(
        doc,
        "The Service is intended for eligible members of the University of Waterloo student community. Trovun may change eligibility requirements when reasonably necessary to protect the marketplace or comply with law.",
    )

    add_heading(doc, "Account responsibilities", 2)
    add_para(
        doc,
        "You are responsible for activity under your account and for maintaining accurate information. You must use reasonable security, protect authentication credentials, and notify Trovun of suspected unauthorized access. Trovun may require re verification or restrict an account when identity, eligibility, or security cannot be confirmed.",
    )

    add_heading(doc, "User transactions", 2)
    add_para(
        doc,
        "A transaction is made directly between the buyer and seller. Each user is responsible for evaluating the other user, item, price, payment method, inspection, pickup, legal compliance, taxes, and any agreed warranty or return term. Unless Trovun expressly states otherwise for a specific feature, Trovun is not a party to the transaction and does not acquire title to any item.",
    )
    add_para(
        doc,
        "Business sellers remain responsible for laws that apply to their activities, including consumer protection, product safety, tax, licensing, disclosure, and record keeping requirements. Nothing in these Terms limits a right or remedy that cannot lawfully be limited.",
    )

    add_heading(doc, "User content and licence", 2)
    add_para(
        doc,
        "You retain ownership of your Content. You grant Trovun a non exclusive, worldwide, royalty free licence to host, store, reproduce, format, display, and transmit that Content only as reasonably necessary to operate, secure, improve, moderate, and promote the Service. This licence ends when the Content is deleted, except for copies retained in backups, transaction records, messages, moderation records, or as required for legal, security, or dispute purposes.",
    )
    add_para(
        doc,
        "You represent that you have the rights needed to submit the Content and that its collection, use, display, and transfer through Trovun will not violate law or another person’s rights. Do not upload confidential material or personal information that is unnecessary for the transaction.",
    )

    add_heading(doc, "Trovun intellectual property", 2)
    add_para(
        doc,
        "The Service, its software, visual design, brand elements, and original materials are owned by Trovun or its licensors and are protected by applicable intellectual property law. These Terms permit personal use of the Service; they do not transfer ownership or authorize copying, scraping, resale, reverse engineering, or use of Trovun branding in a way that suggests endorsement.",
    )

    add_heading(doc, "Service communications", 2)
    add_para(
        doc,
        "Trovun may send communications needed to operate or secure the Service, including verification, password, account, safety, policy, moderation, and transaction related notices. Optional message reminder emails are governed by the Notification Policy and may be disabled in profile settings. Marketing communications, if introduced, will require the consent and unsubscribe controls required by applicable law.",
    )

    add_heading(doc, "Moderation and enforcement", 2)
    add_para(
        doc,
        "Trovun may investigate reported or detected activity and may reject, limit, hide, archive, remove, or preserve Content; restrict marketplace features; suspend or terminate accounts; and refer credible threats or unlawful conduct to appropriate authorities. Trovun may act without advance notice when necessary to address security, legal, fraud, or safety risk. Moderation decisions may rely on reports, account history, technical signals, messages made available for review, and other relevant evidence.",
    )

    add_heading(doc, "Availability and changes", 2)
    add_para(
        doc,
        "Trovun may modify, suspend, or discontinue any feature. The Service may be unavailable because of maintenance, provider outages, security events, or circumstances outside Trovun’s reasonable control. Trovun does not promise uninterrupted or error free operation. Material policy changes will be communicated through the Service or another reasonable method and will apply from the stated effective date.",
    )

    add_heading(doc, "Disclaimers", 2)
    add_para(
        doc,
        "To the maximum extent permitted by law, the Service is provided on an as is and as available basis. Trovun does not warrant the identity, conduct, authority, title, quality, legality, safety, authenticity, availability, price, description, or fitness of any user, listing, or item. Verification of a Waterloo email address confirms control of that address at the relevant time; it is not a guarantee of identity, current enrolment, character, or transaction reliability.",
    )

    add_heading(doc, "Limits of liability", 2)
    add_para(
        doc,
        "To the maximum extent permitted by law, Trovun and its operators, contributors, and service providers will not be liable for indirect, incidental, special, consequential, exemplary, or punitive damages, or for lost profits, data, opportunity, goodwill, or property arising from the Service or a user transaction. Any aggregate liability that cannot lawfully be excluded should be limited to the greater of the amount the user paid directly to Trovun for the Service during the twelve months before the claim and one hundred Canadian dollars. This limitation does not apply where prohibited by law.",
    )

    add_heading(doc, "Indemnity", 2)
    add_para(
        doc,
        "To the extent permitted by law, you agree to indemnify and hold harmless Trovun and its operators from third party claims, losses, and reasonable costs arising from your Content, your transaction, your breach of these policies, or your violation of law or another person’s rights. Trovun will provide reasonable notice of a covered claim and may control its defence and settlement.",
    )

    add_heading(doc, "Suspension termination and survival", 2)
    add_para(
        doc,
        "You may stop using Trovun at any time. Trovun may suspend or terminate access for a material or repeated breach, unlawful conduct, fraud, safety risk, ineligibility, prolonged inactivity, or operational closure. Provisions that by their nature should continue, including ownership, transaction responsibility, disclaimers, liability limits, indemnity, records, and dispute terms, survive termination.",
    )

    add_heading(doc, "Governing law and disputes", 2)
    add_para(
        doc,
        "These Terms are governed by the laws of Ontario and the federal laws of Canada that apply there, without regard to conflict of law rules. Subject to any mandatory consumer right, the courts located in Ontario will have jurisdiction. Before starting a formal claim, users are encouraged to provide written notice describing the issue and requested resolution so the parties can try to resolve it informally.",
    )

    add_heading(doc, "General terms", 2)
    add_para(
        doc,
        "If a provision is unenforceable, it will be limited or removed only to the extent necessary and the remaining provisions will continue. A delay in enforcing a provision is not a waiver. You may not transfer your account or these Terms without Trovun’s consent. Trovun may assign these Terms as part of a reorganization, financing, or transfer of the Service, subject to applicable privacy and consumer protection law. These policies and any feature specific terms form the entire agreement about the Service.",
    )

    # Privacy
    # Let the privacy section follow naturally so the Terms do not leave a
    # nearly empty carry-over page before a forced break.
    add_heading(doc, "6  Privacy Statement", 1)
    add_heading(doc, "Scope and privacy principles", 2)
    add_para(
        doc,
        "This Privacy Statement explains how Trovun collects, uses, discloses, retains, and protects personal information through the Service. Trovun intends to follow applicable Canadian privacy law and the fair information principles of accountability, identified purposes, consent, limited collection, limited use and retention, accuracy, safeguards, openness, access, and recourse.",
    )

    add_heading(doc, "Information Trovun collects", 2)
    privacy_rows = [
        ("Account and verification", "Email address, account identifier, email verification status, authentication records, role, and account timestamps.", "Create and secure accounts, confirm eligibility, prevent abuse, and provide account recovery."),
        ("Student profile", "Name, program, academic year, university, broad residence information provided during onboarding, and optional profile photo.", "Personalize the account, support trust, and show only approved marketplace safe profile fields to eligible users."),
        ("Listings", "Title, description, price, category, condition, offer preference, images, listing status, timestamps, and public meetup address or map coordinates.", "Create, search, display, manage, moderate, and preserve transaction context for listings."),
        ("Messages", "Conversation participants, listing context, message text, read timestamps, delivery metadata, and conversation activity.", "Provide private messaging, unread counts, safety review, and optional message notifications."),
        ("Ratings", "Buyer, seller, listing, conversation, score, and timestamps.", "Enable transaction based seller feedback, prevent duplicate ratings, and calculate aggregate seller ratings."),
        ("Notifications", "Message email preference, recipient address, notification job status, attempts, timing, and provider response metadata.", "Send unread message reminders, honour opt outs, suppress read messages, and diagnose delivery failures."),
        ("Moderation and support", "Reports, listing snapshots, removal reasons, moderator actions, correspondence, and records relevant to safety or disputes.", "Enforce policies, respond to concerns, document decisions, and meet legal obligations."),
        ("Technical and session", "Authentication cookies, session policy and activity timestamps, IP address and device or browser information that may be processed by infrastructure providers, request logs, and security signals.", "Keep users signed in, expire inactive sessions, operate the Service, detect misuse, troubleshoot, and protect security."),
    ]
    add_policy_table(doc, ["Category", "Examples", "Main purposes"], privacy_rows, [1.3, 2.7, 2.7])

    add_heading(doc, "Information other users can see", 2)
    add_para(
        doc,
        "Eligible users may see marketplace safe information such as a shortened display name, program, academic year, university, account age, optional profile photo, aggregate rating, active listings, listing photos, public meetup details, and messages sent directly to them. Trovun does not expose private residence information or account email addresses as marketplace contact information.",
    )
    add_para(
        doc,
        "Messages are visible to the conversation participants and may be accessed by authorized personnel or service providers when reasonably necessary for support, moderation, security, legal compliance, or protection of users. Do not place sensitive information in a listing or message unless it is necessary and you are comfortable sharing it with the intended recipient.",
    )

    add_heading(doc, "How Trovun uses information", 2)
    add_bullets(
        doc,
        [
            "Provide, personalize, maintain, and improve the Service.",
            "Verify eligibility, authenticate users, recover accounts, and maintain session security.",
            "Publish and search listings, display marketplace safe profiles, and facilitate messages and meetups.",
            "Send operational notices and user controlled unread message reminders.",
            "Calculate seller ratings and preserve transaction context after a listing is sold, archived, or deleted.",
            "Detect fraud, spam, prohibited items, security threats, and policy violations.",
            "Investigate reports, enforce policies, protect users, and establish or defend legal claims.",
            "Comply with lawful requests and legal obligations.",
            "Create aggregated or de identified information that does not reasonably identify an individual.",
        ],
    )

    add_heading(doc, "Consent and legal authority", 2)
    add_para(
        doc,
        "Trovun seeks meaningful consent where required and collects information for purposes a reasonable person would consider appropriate. Some processing is necessary to provide a requested service, secure the platform, complete a user initiated communication, enforce the agreement, or comply with law. Optional message email reminders can be disabled. Withdrawing consent may prevent Trovun from providing a feature that depends on the information.",
    )

    add_heading(doc, "When information is disclosed", 2)
    add_bullets(
        doc,
        [
            ("Other users  ", "Marketplace safe profile and listing information is disclosed as described above; messages are disclosed to their participants."),
            ("Service providers  ", "Trovun uses providers for database, authentication, file storage, hosting, email delivery, monitoring, and related infrastructure. Current implementation includes Supabase for core backend services and Resend for notification email delivery."),
            ("Safety and legal compliance  ", "Information may be disclosed when reasonably necessary to respond to lawful process, investigate fraud or a credible safety threat, protect rights and security, or comply with law."),
            ("Organizational change  ", "Information may be transferred in a financing, reorganization, merger, acquisition, or sale of the Service, subject to confidentiality and applicable law."),
            ("With direction or consent  ", "Trovun may disclose information when a user directs it or gives valid consent."),
        ],
    )
    add_para(doc, "Trovun does not sell personal information for money.", bold_lead="Trovun does not sell personal information for money.")

    add_heading(doc, "Cookies and session technology", 2)
    add_para(
        doc,
        "Trovun uses authentication and session cookies required for sign in, security, and continuity. The current web session policy records a policy version and recent activity timestamp and expires an inactive web session after approximately three days. Authentication providers may use additional cookies needed to maintain a secure session. Because these technologies are necessary to deliver the signed in Service, disabling them may prevent account features from working.",
    )

    add_heading(doc, "Email notifications", 2)
    add_para(
        doc,
        "Unread message reminders are enabled by default in the current implementation and may be disabled in profile settings. Delivery waits before sending, suppresses a reminder if the message has already been read, combines rapid activity, and limits reminders for a conversation. The notification email links to the conversation and does not include message content. A delivery already in progress may finish after the preference is changed.",
    )

    add_heading(doc, "Retention and deletion", 2)
    add_para(
        doc,
        "Trovun retains information only as long as reasonably needed for the purposes described in this statement, including account operation, transaction continuity, safety, dispute handling, security, backup recovery, and legal obligations. Retention depends on the record and account state. For example, a conversation may retain a listing title and image reference after the listing is sold or deleted so participants still understand the transaction context. Moderation records may outlast the listing they document.",
    )
    add_para(
        doc,
        "When information is no longer required, Trovun will delete, anonymize, or securely isolate it, subject to backup cycles and legal exceptions. A user may request account or personal information deletion, but Trovun may retain limited records needed for security, fraud prevention, legal compliance, dispute resolution, or the rights of other users.",
    )

    add_heading(doc, "Storage and international processing", 2)
    add_para(
        doc,
        "Trovun and its service providers may process or store information outside Ontario or Canada. Information in another jurisdiction may be accessible to courts, law enforcement, or regulators under that jurisdiction’s laws. Trovun should maintain contractual and technical safeguards appropriate to the sensitivity of the information.",
    )

    add_heading(doc, "Security", 2)
    add_para(
        doc,
        "Trovun uses administrative, technical, and access controls intended to protect information, including verified accounts, private storage, row level database access rules, limited public profile views, signed file access, session expiry, and restricted notification queues. No system is completely secure. Users should use a unique password, protect verification codes, sign out of shared devices, and report suspected compromise promptly.",
    )

    add_heading(doc, "Access correction and complaints", 2)
    add_para(
        doc,
        "Subject to lawful limits, users may ask whether Trovun holds personal information about them, request access, ask for a correction, withdraw optional consent, or raise a privacy concern. Trovun may need to verify identity before responding and may withhold information that would reveal another person’s information, compromise security, or be protected by law. Users may also have the right to complain to the Office of the Privacy Commissioner of Canada or another applicable regulator.",
    )

    add_heading(doc, "Young users", 2)
    add_para(
        doc,
        "The Service is intended for eligible post secondary students who can legally agree to these policies. Trovun does not knowingly design the Service for children. If Trovun learns that an ineligible minor has created an account or submitted personal information, it may restrict the account and delete information as permitted by law.",
    )

    add_heading(doc, "Privacy contact", 2)
    add_para(
        doc,
        "Privacy questions, access requests, correction requests, deletion requests, and complaints should be directed to the privacy contact identified in the Service. Before this statement is published, Trovun must identify its accountable privacy lead, monitored privacy email address, legal operator name, and mailing address.",
    )

    # Notifications
    add_heading(doc, "7  Notification Policy", 1, page_break=True)
    add_heading(doc, "Types of notification", 2)
    notification_rows = [
        ("Account and security", "Verification, password reset, recovery, suspicious activity, eligibility, or material account changes.", "Required when needed to provide or secure the account."),
        ("Marketplace operations", "Important listing, moderation, safety, policy, or service notices.", "Required when reasonably necessary to operate the Service or protect users."),
        ("Unread message reminders", "Email reminder that links to an unread conversation without reproducing message content.", "User controlled in profile settings; currently enabled by default."),
        ("Marketing", "Product announcements, promotions, surveys, or partner messages that are commercial in nature.", "Not part of the current documented workflow. Trovun must obtain and record any consent required before introducing them."),
    ]
    add_policy_table(doc, ["Type", "Examples", "Control"], notification_rows, [1.45, 3.25, 2.0])

    add_heading(doc, "Unread message reminder behaviour", 2)
    add_bullets(
        doc,
        [
            "A new message creates a private notification job for the other participant.",
            "Delivery waits at least one minute and rechecks whether the message remains unread and whether the recipient still permits message emails.",
            "Rapid messages may be combined, and one conversation sends no more than one reminder in a ten minute period.",
            "The email identifies Trovun and links to the conversation but does not include message text.",
            "Delivery failures do not prevent the original in service message from being sent.",
            "A user may disable message emails in profile settings. A reminder already being processed may still arrive.",
        ],
    )

    add_heading(doc, "Commercial messages", 2)
    add_para(
        doc,
        "If Trovun introduces commercial electronic messages, it will use the consent basis permitted by Canada’s Anti Spam Legislation, identify the sender, provide a readily performed unsubscribe mechanism, keep appropriate consent records, and honour unsubscribe requests within the legally required period. Opting out of marketing will not prevent essential account, security, or transaction communications where those messages are otherwise permitted.",
    )

    # Moderation
    add_heading(doc, "8  Reporting moderation and appeals", 1, page_break=True)
    add_heading(doc, "When to report", 2)
    add_bullets(
        doc,
        [
            "A prohibited, recalled, unsafe, stolen, counterfeit, or materially misleading item.",
            "Fraud, impersonation, payment manipulation, account compromise, or suspicious links.",
            "Harassment, threats, discrimination, stalking, sexual exploitation, or disclosure of personal information.",
            "A credible risk of immediate physical harm or other urgent safety concern.",
            "Rating manipulation, academic misconduct, spam, or repeated attempts to evade enforcement.",
        ],
    )

    add_heading(doc, "How Trovun may respond", 2)
    add_para(
        doc,
        "Trovun may ask for additional information, preserve relevant records, remove a listing, restrict messaging, limit account access, suspend or terminate an account, warn affected users, or refer a matter to an appropriate authority. Trovun may prioritize urgent safety issues and is not required to disclose confidential investigation details or information about another user.",
    )

    add_heading(doc, "Appeals", 2)
    add_para(
        doc,
        "A user may request review of a moderation decision through the support method identified in the Service. The request should identify the account and decision, explain why the decision should change, and provide relevant evidence. Trovun may refuse repetitive, abusive, or unsupported requests and may maintain restrictions while a review is pending.",
    )

    add_heading(doc, "Emergencies and law enforcement", 2)
    add_para(
        doc,
        "For an immediate threat, call 911 or the appropriate emergency service. Trovun support is not an emergency service. Requests from law enforcement or other authorities must use the legal contact method identified by Trovun and include valid legal authority and sufficient account or transaction identifiers.",
    )

    # References
    add_heading(doc, "9  Legal references and publication checklist", 1, page_break=True)
    add_heading(doc, "Official reference material", 2)
    references = [
        ("Office of the Privacy Commissioner of Canada  PIPEDA fair information principles", "https://www.priv.gc.ca/en/privacy-topics/privacy-laws-in-canada/the-personal-information-protection-and-electronic-documents-act-pipeda/p_principle/"),
        ("Office of the Privacy Commissioner of Canada  Privacy Guide for Businesses", "https://www.priv.gc.ca/media/2038/guide_org_e.pdf"),
        ("CRTC  Canada Anti Spam Legislation frequently asked questions", "https://www.crtc.gc.ca/eng/com500/faq500.htm"),
        ("Ontario  Advice on shopping online or over the phone", "https://www.ontario.ca/page/advice-shopping-online-or-over-phone"),
        ("Ontario  Consumer Protection Act 2002", "https://www.ontario.ca/laws/statute/02c30"),
        ("Health Canada  Online shopping and product safety", "https://www.canada.ca/en/health-canada/services/buying-consumer-products-online.html"),
        ("Health Canada  Consumer product safety information for direct sellers", "https://www.canada.ca/en/health-canada/services/consumer-product-safety/reports-publications/industry-professionals/canada-consumer-product-safety-act-information-direct-sellers.html"),
        ("Ontario  Rules for selling tobacco and vapour products", "https://www.ontario.ca/page/rules-selling-tobacco-and-vapour-products"),
        ("Ontario  Buying recreational cannabis", "https://www.ontario.ca/page/buying-recreational-cannabis"),
        ("Justice Laws Website  Controlled Drugs and Substances Act", "https://laws-lois.justice.gc.ca/eng/acts/C-38.8/"),
    ]
    for label, url in references:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)
        set_keep(p, together=True)
    add_para(doc, "References reviewed 21 September 2026. Laws and regulatory guidance change. Trovun should review this document periodically and after any material change to the Service.")

    add_heading(doc, "Required decisions before publication", 2)
    checklist_rows = [
        ("Legal operator", "Insert the full legal name of the person or entity operating Trovun and confirm authority to accept these Terms."),
        ("Contact details", "Establish and monitor support, privacy, and legal contact channels and add a service address where required."),
        ("Age and eligibility", "Confirm the minimum age and enrolment or alumni eligibility rules and align registration controls."),
        ("Retention schedule", "Approve record specific retention periods for accounts, listings, images, messages, ratings, notification jobs, moderation records, and backups."),
        ("Service providers", "Confirm current hosting, database, authentication, storage, monitoring, and email providers and execute appropriate data protection terms."),
        ("Incident response", "Assign a privacy lead and approve security incident, breach assessment, notification, and evidence preservation procedures."),
        ("Reporting workflow", "Publish a usable support and reporting path for prohibited items, harassment, fraud, safety concerns, privacy requests, and appeals."),
        ("Consent records", "Confirm default settings and records for notifications, and obtain express consent before any new marketing program when required."),
        ("Insurance and liability", "Have Canadian counsel review disclaimers, liability limits, indemnity, governing law, and insurance needs."),
        ("Accessibility and language", "Review the final web presentation for accessibility and determine whether a French version or other localized notice is required."),
    ]
    add_policy_table(doc, ["Decision", "Publication requirement"], checklist_rows, [1.7, 5.0])

    add_heading(doc, "Document control", 2)
    add_policy_table(
        doc,
        ["Field", "Value"],
        [
            ("Document", "Trovun Marketplace Policies and Community Standards"),
            ("Version", "1.0 publication draft"),
            ("Effective date", "21 September 2026"),
            ("Owner", "Trovun operator to be confirmed"),
            ("Review cycle", "At least annually and after any material product, legal, provider, or data practice change"),
            ("Approval", "Canadian legal and privacy review required before publication"),
        ],
        [1.55, 5.15],
    )

    # Metadata and consistency settings.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            if not run.font.name:
                run.font.name = "Aptos"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
